"""Rendering (separate from core). HTML (phone-friendly), Markdown, XLSX, CSV, JSON.
The app always asks for ONE language (lang="en" or "es"): every heading, label, table header, note and sheet name
comes out in that language only. lang="both" (the old EN/ES toggle page) is kept only for the command-line tools."""
import json, re, html as H, datetime as dt, pathlib, contextvars, functools
import pandas as pd
from .i18n import t, COST

_LANG = contextvars.ContextVar("bellazu_render_lang", default="both")


def _lang():
    return _LANG.get()


def with_lang(fn):
    """Run a top-level renderer with its `lang` argument (positional or keyword) as the current language for bi()/bk()."""
    import inspect
    sig = inspect.signature(fn)

    @functools.wraps(fn)
    def wrap(*a, **k):
        ba = sig.bind(*a, **k); ba.apply_defaults()
        lang = ba.arguments.get("lang", "both")
        tok = _LANG.set(lang if lang in ("en", "es") else "both")
        try:
            return fn(*ba.args, **ba.kwargs)
        finally:
            _LANG.reset(tok)
    return wrap

CSS = """@import url('https://fonts.googleapis.com/css2?family=League+Gothic&family=Inter:wght@300;400;500;600;700&family=Instrument+Serif&display=swap');
:root{--ink:#141414;--mute:#6F6F6F;--line:#E6E0E2;--rose:#D8567C;--rose2:#FBE7ED;--paper:#FFFFFF;--good:#1E7A4B;--bad:#B3264F}
body{font-family:Inter,-apple-system,Segoe UI,Roboto,Arial,sans-serif;max-width:860px;margin:0 auto;padding:18px 18px 40px;color:var(--ink);font-size:15px;line-height:1.55;background:var(--paper)}
.brand{font-family:'Instrument Serif',Georgia,serif;font-size:30px;line-height:1;padding:6px 0 12px;border-bottom:1px solid var(--ink);margin-bottom:18px}.brand i{font-style:normal;color:var(--rose)}
h1{font-family:'League Gothic',Impact,sans-serif;font-weight:400;text-transform:uppercase;font-size:46px;line-height:.95;letter-spacing:.01em;margin:.1em 0 .25em}
h2{font-family:'League Gothic',Impact,sans-serif;font-weight:400;text-transform:uppercase;font-size:28px;line-height:1;letter-spacing:.01em;margin:1.4em 0 .5em;padding-top:.7em;border-top:1px solid var(--line)}
ul{padding-left:1.1em}li{margin:.35em 0}
table{border-collapse:collapse;width:100%;margin:.5em 0 .8em;font-size:.9em}td,th{border-bottom:1px solid var(--line);padding:8px 8px;text-align:left;vertical-align:top}
th{font-size:.72em;letter-spacing:.08em;text-transform:uppercase;color:var(--mute);font-weight:500;border-bottom:1px solid var(--ink)}.num{text-align:right;font-variant-numeric:tabular-nums}
.neg{color:var(--bad);font-weight:600}.pos{color:var(--good);font-weight:600}
.box{background:var(--rose2);border:1px solid #F3C6D3;border-radius:18px;padding:12px 16px;margin:.8em 0}.bad{background:#FDEEF2;border-color:#F0B8C8}.ok{background:#EAF6EF;border-color:#BFE3CD}
.small{font-size:.8em;color:var(--mute)}
.toggle{position:sticky;top:0;background:var(--paper);padding:8px 0;z-index:9;display:flex;gap:6px}.toggle button{padding:8px 14px;border:1px solid var(--ink);background:var(--paper);color:var(--ink);border-radius:100px;
 text-transform:uppercase;letter-spacing:.06em;font-size:.72em;font-family:inherit;cursor:pointer}.toggle button:hover{background:var(--ink);color:var(--paper)}
body.only-en .es{display:none}body.only-es .en{display:none}body.both div.es{display:block;color:var(--mute);margin-top:2px}body.both span.es{display:inline}
body.both span.es:before{content:" / "}a{color:var(--rose);word-break:break-word}
@media(max-width:600px){body{padding:14px}h1{font-size:38px}table{font-size:.8em}td,th{padding:6px 4px}}
@media print{.toggle{display:none}body{max-width:none}}"""
JS = "<script>function L(m){document.body.className=m;try{localStorage.bz=m}catch(e){}};try{if(localStorage.bz)L(localStorage.bz)}catch(e){}</script>"


PLAIN = [(" → ", ": "), ("≈ ", "about "), ("≈", "about "), ("~$", "about $"), (" ~", " about "), ("ASSUMPTIONS", "Estimates"), ("ASSUMPTION", "Estimate"),
         ("SUPUESTOS", "Estimados"), ("SUPUESTO", "Estimado"), ("PROXY", "Rough guide"), (" — ", ": "), (" -> ", " to "), ("→", "to"),
         ("⚠️", ""), ("NOT ", "not "), ("Arbitrage Scan", "Rental Finder"), ("Análisis de Arbitraje", "Buscador de Alquileres"),
         ("(edit assumptions.yaml)", ""), ("(editar assumptions.yaml)", ""), ("edit assumptions.yaml", "our standard estimates"),
         ("Principal & interest", "Mortgage payment (loan + interest)"), ("MIP/PMI", "required with a small down payment"),
         ("Furniture (amortized)", "Furniture (spread over 3 years)"), ("Muebles (amortizados)", "Muebles (repartidos en 3 años)"),
         ("NET per month", "Left over each month"), ("NETO por mes", "Queda cada mes"), ("STR insurance", "Short-stay insurance"),
         ("Short-term rental comps (Inside Airbnb)", "Nearby Airbnb listings (Inside Airbnb)"), ("Room-rent comps", "Room rents nearby"),
         ("Comparables de alquiler a corto plazo", "Anuncios de Airbnb cercanos"), ("Comparables de renta por cuarto", "Rentas de cuartos cercanas"),
         ("Long-term rent estimate", "What a tenant would pay"), ("Reference rents", "Reference rents"), ("Referencias", "Rentas de referencia"),
         ("STR profit/mo", "Airbnb profit/mo"), ("STR rev/yr", "Airbnb income/yr"), ("STR profit", "Airbnb profit"), ("30+night", "30+ night"),
         ("Rentals ranked by estimated STR profit", "Rentals ranked by estimated profit"), ("ganancia STR estimada", "ganancia estimada"),
         ("comps", "similar listings"), ("Comps", "Similar listings"), ("LTR", "long-term rent"), (" STR ", " Airbnb-style "), ("(STR)", "(Airbnb-style)"),
         ("Median rent", "Typical rent"), ("median", "typical"), ("mediana", "típica"), ("range (25th-75th pct)", "usual range"), ("rango (percentil 25-75)", "rango usual"),
         ("Ownership type", "Type of home"), ("Tipo de propiedad", "Tipo de vivienda"), ("Cash to close (est.)", "Cash to close (estimate)"),
         ("Efectivo para cerrar (est.)", "Dinero para cerrar (estimado)"), ("unknown / desconocido", "not found / no encontrado")]


PLAIN_ES = [("≈ ", "aprox. "), ("≈", "aprox. "), ("~$", "aprox. $"), (" ~", " aprox. "), ("SUPUESTOS", "Estimados"), ("SUPUESTO", "Estimado"),
            (" → ", ": "), (" — ", ": "), (" -> ", " a "), ("→", "a"), ("5-year", "5 años"), (" via ", " vía "), ("week of", "semana del"), ("⚠️", ""), ("NO ", "no "), ("Análisis de Arbitraje", "Buscador de Alquileres"),
            ("(editar assumptions.yaml)", ""), ("Capital e intereses", "Pago de la hipoteca (préstamo + interés)"), ("MIP/PMI", "obligatorio con poco enganche"),
            ("Muebles (amortizados)", "Muebles (repartidos en 3 años)"), ("NETO por mes", "Queda cada mes"), ("Seguro STR", "Seguro de estadías cortas"),
            ("Comparables de alquiler a corto plazo", "Anuncios de Airbnb cercanos"), ("Comparables de renta por cuarto", "Rentas de cuartos cercanas"),
            ("Referencias", "Rentas de referencia"), ("ganancia STR estimada", "ganancia estimada"), ("Ganancia STR", "Ganancia Airbnb"), ("ingreso STR", "ingreso Airbnb"),
            (" STR ", " tipo Airbnb "), ("mediana", "típica"), ("rango (percentil 25-75)", "rango usual"), ("Tipo de propiedad", "Tipo de vivienda"),
            ("Efectivo para cerrar (est.)", "Dinero para cerrar (estimado)"), ("efectivo para cerrar", "dinero para cerrar"), ("desconocido", "no encontrado")]


def _plain_text(t, lang="en"):
    for a, b in (PLAIN_ES if lang == "es" else PLAIN):
        t = t.replace(a, b)
    return re.sub(r"[ \t]{2,}", " ", t)


def plainify(doc, lang="both"):
    """Plain-language pass over the visible text of a report (never touches tags, attributes, URLs or numbers)."""
    if lang in ("en", "es"):
        head, sep, body = doc.partition("<body")
        body = re.sub(r">([^<]+)<", lambda m: ">" + _plain_text(m.group(1), lang) + "<", body)
        head = re.sub(r"<title>(.*?)</title>", lambda m: "<title>" + _plain_text(m.group(1), lang) + "</title>", head, flags=re.S)
        return head + sep + body
    head, sep, body = doc.partition("<body")
    es_fix = lambda t: t.replace("≈ ", "aprox. ").replace("≈", "aprox. ").replace("~$", "aprox. $").replace(" ~", " aprox. ")
    body = re.sub(r'(<(\w+) class="es">)(.*?)(</\2>)', lambda m: m.group(1) + re.sub(r">([^<]+)<", lambda n: ">" + es_fix(n.group(1)) + "<", ">" + m.group(3) + "<")[1:-1] + m.group(4), body, flags=re.S)
    body = re.sub(r">([^<]+)<", lambda m: ">" + _plain_text(m.group(1)) + "<", body)
    head = re.sub(r"<title>(.*?)</title>", lambda m: "<title>" + _plain_text(m.group(1)) + "</title>", head, flags=re.S)
    return head + sep + body


def _open(title, lang):
    if lang in ("en", "es"):          # one language: no toggle, no hidden copy of the other language
        return (f"<!doctype html><html lang='{lang}'><head><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'><base target='_blank'>"
                f"<title>{title}</title><style>{CSS}</style></head><body><div class='brand'>Bella<i>Zu</i></div>")
    cls = "both"
    return (f"<!doctype html><html><head><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'><base target='_blank'>"
            f"<title>{title}</title><style>{CSS}</style></head><body class='{cls}'><div class='brand'>Bella<i>Zu</i></div>"
            "<div class='toggle'><button onclick=\"L('only-en')\">English</button><button onclick=\"L('only-es')\">Español</button><button onclick=\"L('both')\">EN + ES</button></div>")


def money(v, dec=0):
    if v is None or (isinstance(v, float) and v != v):
        return "—"
    return f"-${abs(v):,.{dec}f}" if v < 0 else f"${v:,.{dec}f}"


def pct(v):
    return "—" if v is None else f"{v:.0%}"


def slug(s):
    return re.sub(r"[^A-Za-z0-9]+", "_", s).strip("_")[:60]


def bi(en, es, tag="span"):
    lg = _lang()
    if lg == "en":
        return en if tag == "span" else f"<{tag}>{en}</{tag}>"
    if lg == "es":
        return es if tag == "span" else f"<{tag}>{es}</{tag}>"
    return f'<{tag} class="en">{en}</{tag}><{tag} class="es">{es}</{tag}>'


def _js():
    return JS if _lang() == "both" else ""


def _mo():
    return "/mes" if _lang() == "es" else "/mo"


def bk(key, tag="span"):
    return bi(H.escape(t(key, "en")), H.escape(t(key, "es")), tag)


def _table(headers, rows, num_cols=()):
    h = "<table><tr>" + "".join(f"<th>{x}</th>" for x in headers) + "</tr>"
    for r in rows:
        h += "<tr>" + "".join(f'<td class="{"num" if i in num_cols else ""}">{c}</td>' for i, c in enumerate(r)) + "</tr>"
    return h + "</table>"


def safe_url(u):
    u = str(u or "")
    return H.escape(u) if u.startswith(("http://", "https://")) else "#"


def _src_link(s):
    if not s:
        return ""
    s = str(s)
    if " ; " in s:
        return " · ".join(_src_link(x.strip()) for x in s.split(" ; "))
    if s.startswith("http"):
        return f'<a href="{safe_url(s)}">{H.escape(s.split("/")[2])}</a>'
    return H.escape(s)


def _rule_src(x):
    m = re.search(r"https?://\S+", x)
    if not m:
        return H.escape(x)
    u = m.group(0).rstrip(")")
    label = x.replace(m.group(0), "").strip(" :()") or u.split("/")[2]
    if _lang() == "es":
        label = {"news": "noticias", "town website": "sitio del municipio"}.get(label.lower(), label)
    return f'<a href="{safe_url(u)}">{H.escape(label[:60])}</a>'


# ------------------------------------------------------------------ property
INC_ES = {"taxes": "impuestos", "utilities": "servicios (luz, gas, calefacción)", "internet": "internet", "heat/hot water": "calefacción/agua caliente"}
FACT_ES = {"single-family": "casa unifamiliar", "two-family": "casa de dos familias", "FOR_SALE": "en venta", "FOR_RENT": "en alquiler", "SOLD": "vendida", "PENDING": "pendiente", "OFF_MARKET": "fuera del mercado",
           "co-op": "co-op (cooperativa)", "condo": "condominio", "house": "casa", "townhouse": "casa adosada", "multi-family": "multifamiliar"}
SRC_ES = {"parsed from listing description": "tomado de la descripción del anuncio", "user entered": "escrito por usted", "you entered": "escrito por usted", "user input": "escrito por usted",
          "manual": "escrito por usted", "town average": "promedio del pueblo"}
BED_RULE_ES = {"exact": "exacta", "nearest-k proxy": "los más cercanos (aproximación)", "exact (nearest-k proxy)": "exacta (los más cercanos, aproximación)"}


def property_bullets(r):
    f, sc = r["facts"], {s["key"]: s for s in r["scenarios"]}
    out = []
    inc = ", ".join(f.get("hoa_includes") or [])
    inc_es = ", ".join(INC_ES.get(x, x) for x in (f.get("hoa_includes") or []))
    out.append((f"Listed at {money(f.get('price'))}; {f.get('beds') or '?'} bd / {f.get('baths') or '?'} ba {f.get('ownership') or ''}; maintenance/HOA {money(f.get('hoa_monthly'))}/mo" + (f" (includes {inc})." if inc else "."),
                f"Precio {money(f.get('price'))}; {f.get('beds') or '?'} hab / {f.get('baths') or '?'} baños {f.get('ownership') or ''}; mantenimiento/HOA {money(f.get('hoa_monthly'))}/mes" + (f" (incluye {inc_es})." if inc else ".")))
    fh = r.get("fha") or {}
    if fh.get("eligible") is False:
        out.append((fh["note_en"], fh["note_es"]))
    o = sc.get("owner_roommates")
    if o:
        d = o.get("front_end_dti")
        inc = r.get("income_annual")
        if d is not None and inc:
            out.append((f"Living there: total cost {money(o['total_cost'])}/mo, cash to close ≈ {money(o['loan']['cash_to_close_est'])}. Housing payment = {d:.0%} of the {money(inc)}/yr income entered "
                        + ("— far above the ~31-43% lenders usually allow; qualifying is unlikely." if d > 0.45 else ("— high; a lender may push back." if d > 0.36 else "— within typical guidelines.")),
                        f"Viviendo allí: costo total {money(o['total_cost'])}/mes, efectivo para cerrar ≈ {money(o['loan']['cash_to_close_est'])}. El pago de vivienda = {d:.0%} del ingreso ingresado de {money(inc)}/año "
                        + ("— muy por encima del ~31-43% que suelen permitir los prestamistas; es poco probable calificar." if d > 0.45 else ("— alto; el prestamista podría objetar." if d > 0.36 else "— dentro de lo normal."))))
        else:
            out.append((f"Living there: total cost {money(o['total_cost'])}/mo, cash to close ≈ {money(o['loan']['cash_to_close_est'])}. Enter your annual income to see the housing-payment ratio lenders check.",
                        f"Viviendo allí: costo total {money(o['total_cost'])}/mes, efectivo para cerrar ≈ {money(o['loan']['cash_to_close_est'])}. Ingrese su ingreso anual para ver la proporción que revisan los prestamistas."))
        out.append((f"With {o['rooms_rented']} roommate(s) at ~{money(o['room_rent_each'])} each ({o['room_rent_basis']}), your own housing cost ≈ {money(o['own_net_housing_cost'])}/mo.",
                    f"Con {o['rooms_rented']} compañero(s) a ~{money(o['room_rent_each'])} c/u, su costo propio de vivienda ≈ {money(o['own_net_housing_cost'])}/mes."))
    ic = r.get("coop_income_check")
    if ic:
        bi_ = ic.get("buyer_income")
        short = bi_ is not None and bi_ < ic['required_income_with_mortgage']
        out.append((f"Building rule {ic['multiple']}:1 income-to-housing-cost: needs ≈ {money(ic['required_income_maintenance_only'])}/yr on maintenance alone ({money(ic['required_income_with_mortgage'])} with the mortgage)"
                    + (f" vs {money(bi_)} entered." + (" The board would likely reject." if short else "") if bi_ else "."),
                    f"Regla del edificio {ic['multiple']}:1 ingreso/costo: requiere ≈ {money(ic['required_income_maintenance_only'])}/año solo por mantenimiento ({money(ic['required_income_with_mortgage'])} con hipoteca)"
                    + (f" vs {money(bi_)} ingresado." + (" La junta probablemente rechazaría." if short else "") if bi_ else ".")))
    rc = r.get("rent_compare") or {}
    if rc.get("rentcast_ok"):
        g = rc.get("gap_pct")
        out.append((f"RentCast rent estimate {money(rc['rentcast_rent'])}/mo (range {money(rc.get('rentcast_low'))}–{money(rc.get('rentcast_high'))}, {rc['rentcast_n']} comps)"
                     + (f" vs free-source comps {money(rc['free_median'])} (n={rc['free_n']}); gap {g:+.0%}" + (" ⚠️ BIG GAP" if rc.get('big_gap') else "") if g is not None else "")
                     + f". Using: {'RentCast' if rc.get('chosen') == 'rentcast' else ('free comps' if rc.get('chosen') == 'free' else rc.get('chosen'))} (more comps).",
                     f"Estimado de renta RentCast {money(rc['rentcast_rent'])}/mes (rango {money(rc.get('rentcast_low'))}–{money(rc.get('rentcast_high'))}, {rc['rentcast_n']} comparables)"
                     + (f" vs comparables gratuitos {money(rc['free_median'])} (n={rc['free_n']}); diferencia {g:+.0%}" + (" ⚠️ GRAN DIFERENCIA" if rc.get('big_gap') else "") if g is not None else "")
                     + f". Se usa: {'RentCast' if rc.get('chosen') == 'rentcast' else ('comparables gratuitos' if rc.get('chosen') == 'free' else rc.get('chosen'))} (más comparables)."))
    le = r["ltr"]["estimate"]
    inv = sc.get("investment_ltr")
    if inv and rc.get("chosen") == "rentcast":
        out.append((f"Long-term rent used: {money(rc['rentcast_rent'])}/mo (RentCast, {rc['rentcast_n']} comps — more than the {rc.get('free_n', 0)} free-source comps). As a pure rental on an investment loan: {money(inv['net_monthly'])}/mo.",
                    f"Renta a largo plazo usada: {money(rc['rentcast_rent'])}/mes (RentCast, {rc['rentcast_n']} comparables — más que los {rc.get('free_n', 0)} gratuitos). Como alquiler con préstamo de inversión: {money(inv['net_monthly'])}/mes."))
    elif le.get("ok") and inv:
        out.append((f"Long-term rent estimate {money(le['median'])}/mo (range {money(le['p25'])}–{money(le['p75'])}, {le['n']} same-bedroom listings within {le['radius_km']} km). As a pure rental on an investment loan: {money(inv['net_monthly'])}/mo.",
                    f"Renta estimada a largo plazo {money(le['median'])}/mes (rango {money(le['p25'])}–{money(le['p75'])}, {le['n']} anuncios similares a {le['radius_km']} km). Como alquiler con préstamo de inversión: {money(inv['net_monthly'])}/mes."))
    s = sc.get("str")
    if s and not s.get("legal"):
        out.append((f"Airbnb/short-term: NOT legal here → not computed. {r['str_rules'].get('summary_en','')}", f"Airbnb/corto plazo: NO es legal aquí → no calculado. {r['str_rules'].get('summary_es','')}"))
    elif s:
        out.append((f"Short-term rental (rules applied): {money(s['net_monthly'])}/mo.", f"Alquiler a corto plazo (con reglas): {money(s['net_monthly'])}/mes."))
    m = sc.get("mtr")
    if m:
        out.append((f"Legal alternative — furnished 30+ night rental: ≈ {money(m['net_monthly'])}/mo ({m['caveat_en']})", f"Alternativa legal — alquiler amueblado 30+ noches: ≈ {money(m['net_monthly'])}/mes ({m['caveat_es']})"))
    return out


def _km(v):
    try:
        return f"{float(v):.1f}".rstrip("0").rstrip(".")
    except Exception:
        return str(v)


def _rc_status_text(s):
    if _lang() == "es":
        return {"budget": "RentCast no se usó (límite de consultas de esta visita)", "no_key": "RentCast no está configurado", "quota": "Se alcanzó el límite mensual de RentCast; se usaron fuentes gratuitas",
                "not_found": "RentCast no tuvo estimado para esta dirección", None: "no solicitado"}.get(s, f"RentCast no disponible ({s})")
    return {"budget": "RentCast not used (lookup limit for this visit)", "no_key": "RentCast not configured", "quota": "RentCast monthly limit reached — free sources used",
            "not_found": "RentCast had no estimate for this address", None: "not requested"}.get(s, f"RentCast unavailable ({s})")


CMP_CSS = ("<style>.cmp{display:grid;grid-template-columns:repeat(3,minmax(0,1fr));gap:8px;margin:.6em 0}.cmp div.c{border:1px solid var(--line);border-radius:16px;padding:10px 8px}"
           ".cmp .h{font-size:.72em;text-transform:uppercase;letter-spacing:.06em;font-weight:600}.cmp .n{font-family:'League Gothic',Impact,sans-serif;font-size:30px;line-height:1.05}"
           ".cmp .s,.cmp .b{font-size:.75em;color:var(--mute)}.cmp .b{border-top:1px solid var(--line);margin-top:6px;padding-top:6px;color:var(--ink)}</style>")


def compare_html(cv):
    """The 3-column compare strip (built by the app from bellazu.compare.labels) for the downloadable reports."""
    if not cv:
        return ""
    cells = "".join(f"<div class='c'><div class='h'>{bi(H.escape(c['title'][0]), H.escape(c['title'][1]))}</div><div class='s'>{bi(H.escape(c['pay_lbl'][0]), H.escape(c['pay_lbl'][1]))}</div>"
                    f"<div class='n'>{H.escape(c['pay'])}</div><div class='s'>{bi(H.escape(c['sub'][0]), H.escape(c['sub'][1]))}</div>"
                    f"<div class='b'>{bi(H.escape(c['badge'][0]), H.escape(c['badge'][1]))}</div></div>" for c in cv["cols"])
    extra = "".join(f"<li>{bi(H.escape(a), H.escape(b))}</li>" for a, b in ([cv["drive"]] if cv.get("drive") else []) + list(cv.get("lines") or []))
    return (CMP_CSS + f"<h2>{bi('What you pay each month (same FHA loan)', 'Lo que paga al mes (mismo préstamo FHA)')}</h2><div class='cmp'>{cells}</div>"
            + (f"<ul class='small'>{extra}</ul>" if extra else "")
            + "<p class='small'>" + bi("Drive times are typical, not live. Rush-hour range = off-peak route time × 1.68 (TomTom Traffic Index 2025, New York) plus 5-45 min crossing queue. Hudson crossings and Midtown have tolls: panynj.gov/bridges-tunnels/en/tolls.html, congestionreliefzone.mta.info.",
                                      "Los tiempos en carro son típicos, no en vivo. Rango en hora pico = tiempo sin tráfico × 1.68 (TomTom 2025, Nueva York) más 5-45 min de fila en el cruce. Los cruces del Hudson y Midtown tienen peajes: panynj.gov/bridges-tunnels/en/tolls.html, congestionreliefzone.mta.info.") + "</p>")


def _cv_rows(cv, lang="both"):
    if lang in ("en", "es"):
        i = 1 if lang == "es" else 0
        return [{"option": c["title"][i], "label": c["pay_lbl"][i], "per_month": c["pay"], "detail": c["sub"][i], "rules": c["badge"][i]} for c in (cv or {}).get("cols", [])]
    return [{"option_en": c["title"][0], "option_es": c["title"][1], "label": c["pay_lbl"][0], "per_month": c["pay"], "detail": c["sub"][0], "rules": c["badge"][0]} for c in (cv or {}).get("cols", [])]


@with_lang
def property_html(r, lang="both", cv=None):
    f, fs = r["facts"], r["fact_sources"]
    parts = [_open(("Reporte de casa BellaZu" if lang == "es" else "BellaZu Home Report") + f": {H.escape(r['address'])}", lang),
             f"<h1>{bk('prop_title')}</h1><div><b>{H.escape(r['address'])}</b></div><div class='small'>{bk('generated')}: {r['generated']} ET · BellaZu v0.1</div>",
             compare_html(cv),
             f"<h2>{bk('bottom_line')}</h2><ul>" + "".join(f"<li>{bi(H.escape(a), H.escape(b), 'div')}</li>" for a, b in property_bullets(r)) + "</ul>"]
    if r["warnings"]:
        from .simple import plain_warnings
        ws = [bi(H.escape(a), H.escape(b)) for a, b in plain_warnings(r)] if lang in ("en", "es") else [H.escape(w) for w in r["warnings"]]
        parts.append("<div class='box'><b>" + bk("warnings") + "</b><ul>" + "".join(f"<li>{w}</li>" for w in ws) + "</ul></div>")
    rows = []
    for k in ["price", "beds", "baths", "sqft", "hoa_monthly", "taxes_annual", "ownership", "hoa_includes", "status", "year_built"]:
        v = f.get(k)
        if k in ("price", "hoa_monthly", "taxes_annual") and v:
            v = money(float(v))
        if isinstance(v, list):
            v = ", ".join(INC_ES.get(x, x) for x in v) if lang == "es" else ", ".join(v)
        if lang == "es" and isinstance(v, str):
            v = FACT_ES.get(v, v)
        rows.append([bk(k), H.escape(str(v)) if v not in (None, "") else "<i>" + bi("unknown", "desconocido") + "</i>", _src_link(fs.get(k))])
    if lang == "es":
        rows = [[a, b, SRC_ES.get(re.sub(r"<[^>]+>", "", c), c)] for a, b, c in rows]
    parts.append(f"<h2>{bk('facts')}</h2>" + _table([bk("field"), bk("value"), bk("source")], rows))
    bp = r.get("building_policy")
    parts.append(f"<h2>{bk('building_policy')}</h2>" + (f"<p>{bi(H.escape(bp.get('en','')), H.escape(bp.get('es','')))}<br><span class='small'>{_src_link(bp.get('source'))}</span></p>" if bp else
                 "<p>" + bi("Not found in free sources. Ask the listing agent for the co-op/condo house rules: sublet policy, roommate/boarder rules, minimum lease, board approval, flip tax.",
                           "No encontrado en fuentes gratuitas. Pida al agente las reglas de la co-op/condominio: subarriendo, compañeros de cuarto, plazo mínimo, aprobación de la junta, 'flip tax'.") + "</p>"))
    sr = r["str_rules"]
    cls = "bad" if not sr.get("str_legal_for_owner") else "ok"
    parts.append(f"<h2>{bk('legality')}</h2><div class='box {cls}'>{bi(H.escape(sr.get('summary_en','')), H.escape(sr.get('summary_es','')))}<div class='small'>" + bi("Sources", "Fuentes") + ": " + "; ".join(_rule_src(x) for x in sr.get("sources", [])) + f" · {bi('verified', 'verificado')} {sr.get('last_verified')}</div></div>")
    fh = r.get("fha") or {}
    if fh.get("note_en"):
        parts.append(f"<h2>{bk('fha')}</h2><p>{bi(H.escape(fh['note_en']), H.escape(fh['note_es']))}</p>")
    # LTR
    le = r["ltr"]["estimate"]
    parts.append(f"<h2>{bk('ltr')}</h2>")
    if le.get("ok"):
        parts.append(f"<p><b>{money(le['median'])}</b>{_mo()} · {bk('range')}: {money(le['p25'])}–{money(le['p75'])} · n={le['n']} · ≤{le['radius_km']} km</p>")
    rows = [[H.escape(c["source"]), H.escape(str(c.get("title") or ""))[:48], money(c["price"]), int(c["beds"]) if c.get("beds") is not None else "", int(c["sqft"]) if isinstance(c.get("sqft"), (int, float)) and c.get("sqft") == c.get("sqft") else (c.get("sqft") or ""), f"{c['dist_km']:.1f} km", f"<a href='{safe_url(c.get('url'))}'>{bi('link', 'enlace')}</a>"] for c in r["ltr"]["comps"][:10]]
    if rows:
        parts.append(_table([bi("Source", "Fuente"), bi("Listing", "Anuncio"), bi("Rent", "Renta"), bi("Bd", "Hab"), bi("Sqft", "Pies²"), bi("Dist", "Dist."), ""], rows, (2, 3, 4, 5)))
    rc = r.get("rent_compare") or {}
    if rc:
        crow = [[bi("Free-source comps (Craigslist/Rent.com/Redfin)", "Comparables gratuitos (Craigslist/Rent.com/Redfin)"), money(rc.get("free_median")), rc.get("free_n") or 0, "✅" if rc.get("chosen") == "free" else ""],
                [bi("RentCast rent estimate (AVM)", "Estimado de renta RentCast (AVM)"), (money(rc.get("rentcast_rent")) + (f" <span class='small'>({money(rc.get('rentcast_low'))}–{money(rc.get('rentcast_high'))})</span>" if rc.get("rentcast_low") else "")) if rc.get("rentcast_ok") else f"<i>{H.escape(_rc_status_text(rc.get('rentcast_status')))}</i>", rc.get("rentcast_n") or 0, "✅" if rc.get("chosen") == "rentcast" else ""]]
        if rc.get("chosen") == "hud_safmr":
            crow.append([bi("HUD Small Area FMR (fallback)", "HUD Small Area FMR (respaldo)"), money((r["benchmarks"].get("hud_safmr") or {}).get(f"{min(int(f.get('beds') or 2), 4)}br")), "—", "✅"])
        parts.append(_table([bi("Estimate", "Estimado"), bi("Rent/mo", "Renta/mes"), bi("Comps", "Comparables"), bi("Used", "Usado")], crow, (1, 2)))
        if rc.get("big_gap"):
            parts.append(f"<div class='box bad'>{bi(H.escape(rc.get('note_en', '')), H.escape(rc.get('note_es', '')))}</div>")
        parts.append("<p class='small'>" + bi("BellaZu uses the estimate built on more comparable listings.", "BellaZu usa el estimado con más anuncios comparables.") + "</p>")
    rcc = (r.get("rentcast") or {}).get("avm_comps") or []
    if rcc:
        rows = [[H.escape(str(c.get("title") or ""))[:48], money(c.get("price")), int(c["beds"]) if c.get("beds") is not None else "", c.get("sqft") or "", f"{c.get('dist_km') or 0:.1f} km", c.get("days_old") if c.get("days_old") is not None else "", f"<a href='{safe_url(c.get('url'))}'>{bi('map', 'mapa')}</a>"] for c in rcc[:10]]
        parts.append("<p><b>" + bi("RentCast comparables", "Comparables de RentCast") + "</b></p>" + _table([bi("Address", "Dirección"), bi("Rent", "Renta"), bi("Bd", "Hab"), bi("Sqft", "Pies²"), bi("Dist", "Dist."), bi("Days old", "Días"), ""], rows, (1, 2, 3, 4, 5)))
    b = r["benchmarks"]
    beds = min(int(f.get("beds") or 2), 4)
    bl = []
    if b.get("hud_safmr"):
        bl.append(bi(f"HUD FY2026 Small Area FMR ZIP {b['hud_safmr']['zip']}, {beds} bd: ", f"HUD FY2026 Small Area FMR código postal {b['hud_safmr']['zip']}, {beds} hab: ")
                  + f"<b>{money(b['hud_safmr'][f'{beds}br'])}</b> " + bi("(40th percentile gross rent)", "(renta bruta, percentil 40)"))
    if b.get("acs_median_gross_rent"):
        a = b["acs_median_gross_rent"]
        bl.append(bi(f"Census {a.get('release')} median gross rent ZIP {r['zip']}, {beds} bd: ", f"Censo {a.get('release')} renta bruta mediana código postal {r['zip']}, {beds} hab: ")
                  + f"<b>{money(a.get(f'{beds}br'))}</b> " + bi("(all tenants, lags the market)", "(todos los inquilinos, va atrasado frente al mercado)"))
    if b.get("rentcast_avm"):
        bl.append(f"RentCast AVM: <b>{money(b['rentcast_avm'].get('rent'))}</b>")
    parts.append(f"<p class='small'><b>{bk('benchmarks')}:</b> " + " · ".join(bl) + "</p>")
    ro = r["rooms"]
    if ro.get("ok"):
        parts.append(f"<h2>{bk('rooms')}</h2><p>{money(ro['median'])}{_mo()} {bk('median')} · {money(ro['p25'])}–{money(ro['p75'])} · n={ro['n']} ({bi('Craigslist rooms & shares', 'cuartos y casas compartidas de Craigslist')})"
                     + (f" <span class='small'>{H.escape(ro.get('note',''))}</span>" if lang != "es" else "") + "</p>")
    # STR
    st = r["str"]
    parts.append(f"<h2>{bk('str')}</h2>")
    if st.get("ok"):
        s = st["summary"]
        parts.append(f"<p>n={s['n']} {bi('active entire-home listings', 'anuncios activos de casa completa')} ≤{_km(st['radius_km'])} km ({bi('bedrooms rule', 'regla de habitaciones')}: {H.escape(BED_RULE_ES.get(str(st['bed_rule']), str(st['bed_rule'])) if lang == 'es' else str(st['bed_rule']))}) · {bi('data', 'datos')}: {', '.join(st['datasets'])}<br>"
                     f"{bi('Nightly (listed)', 'Precio por noche (anunciado)')} {bk('median')}: <b>{money(s['adr_median'])}</b> · {bi('occupancy (SF model)', 'ocupación (modelo SF)')}: <b>{pct(s['occ_median_sf_model'])}</b> · {bi('income/yr', 'ingreso/año')}: <b>{money(s['revenue_median'])}</b> ({money(s['revenue_p25'])}–{money(s['revenue_p75'])})</p>")
        rows = [[H.escape(str(c.get("neighbourhood_cleansed"))), int(c["bedrooms"]) if c.get("bedrooms") is not None else "", money(c.get("price_num")), c.get("estimated_occupancy_l365d"), money(c.get("estimated_revenue_l365d")), f"{c['dist_km']:.1f}", f"<a href='{safe_url(c.get('listing_url'))}'>airbnb</a>"] for c in st["comps"][:8]]
        parts.append(_table([bi("Area", "Zona"), bi("Bd", "Hab"), bi("Nightly", "Por noche"), bi("Nights/yr", "Noches/año"), bi("Income/yr", "Ingreso/año"), "km", ""], rows, (1, 2, 3, 4, 5)))
        parts.append("<p class='small'>" + bi("Method: Inside Airbnb 'San Francisco model' — booked nights = reviews in last 12 months ÷ 50% review rate × max(3, minimum nights), capped at 70%. Price is the listed nightly price for one quoted date window (not realized ADR), excludes cleaning fees.",
                                              "Método: modelo 'San Francisco' de Inside Airbnb — noches reservadas = reseñas de 12 meses ÷ 50% × max(3, noches mínimas), tope 70%. El precio es el precio listado para una fecha (no el ADR real), sin tarifa de limpieza.") + "</p>")
    # scenarios
    sc = r["scenarios"]
    keys = []
    for s in sc:
        for k in (s.get("costs") or {}):
            if k not in keys:
                keys.append(k)
    head = [""] + [bi(H.escape(s["label_en"]), H.escape(s["label_es"])) for s in sc]
    rows = []
    rows.append([f"<b>{bi('Income', 'Ingreso')}</b>"] + [money(sum((s.get("income") or {}).values())) if s.get("income") else (money(s.get("revenue_monthly")) if s.get("revenue_monthly") else (money(s["revenue_annual"] / 12) if s.get("revenue_annual") else "")) for s in sc])
    for k in keys:
        rows.append([bk(k)] + [("-" + money(s["costs"][k])) if s.get("costs") and s["costs"].get(k) else ("" if s.get("legal") else "") for s in sc])
    rows.append([f"<b>{bk('net')}</b>"] + [f"<span class='{'neg' if s['net_monthly'] < 0 else 'pos'}'>{money(s['net_monthly'])}</span>" if s.get("net_monthly") is not None else f"<i>{bk('not_legal')}</i>" for s in sc])
    rows.append([bk("cash_to_close")] + [money(s["loan"]["cash_to_close_est"]) if s.get("loan") else "" for s in sc])
    rows.append([bi("Rate", "Tasa")] + [f"{s['loan']['rate_pct']:.2f}% · {money(s['loan']['down_payment'])} {bi('down', 'de enganche')}" if s.get("loan") else "" for s in sc])
    parts.append(f"<h2>{bk('scen')}</h2>" + _table(head, rows, tuple(range(1, len(sc) + 1))))
    cav = [bi(H.escape(s.get("caveat_en", "")), H.escape(s.get("caveat_es", ""))) for s in sc if s.get("caveat_en")]
    cav += [bi(H.escape(s.get("dti_note_en", "")), H.escape(s.get("dti_note_es", ""))) for s in sc if s.get("dti_note_en")]
    parts.append("<ul class='small'>" + "".join(f"<li>{c}</li>" for c in cav) + "</ul>")
    parts.append(_assump_html(r))
    parts.append(_sources_html(r))
    parts.append(f"<p class='small'>{bk('disclaimer')}</p>{_js()}</body></html>")
    return plainify("\n".join(parts), lang)


def _assump_html(r):
    A, rt = r["assumptions"], r.get("rates") or {}
    items = []
    fi = A["financing"]
    s = A["str"]
    if rt:
        items.append(bi(f"30-year base rate {rt['base']:.2f}% ({H.escape(rt['source'])}); investment loan +{fi['investment_rate_adjust']:.2f}",
                        f"Tasa base a 30 años {rt['base']:.2f}% ({H.escape(rt['source'])}); préstamo de inversión +{fi['investment_rate_adjust']:.2f}"))
    items.append(bi(f"FHA {fi['fha_down_pct']:.1%} down, upfront MIP {fi['fha_upfront_mip_pct']:.2%}, annual MIP {fi['fha_annual_mip_pct']:.2%}; owner conventional/co-op {fi['owner_conv_down_pct']:.0%} down; investment {fi['investment_down_pct']:.0%} down; closing costs {fi['closing_cost_pct']:.0%}",
                    f"FHA {fi['fha_down_pct']:.1%} de enganche, MIP inicial {fi['fha_upfront_mip_pct']:.2%}, MIP anual {fi['fha_annual_mip_pct']:.2%}; convencional/co-op para vivir {fi['owner_conv_down_pct']:.0%} de enganche; inversión {fi['investment_down_pct']:.0%} de enganche; costos de cierre {fi['closing_cost_pct']:.0%}"))
    cl = H.escape(str(s['cleaning_cost_per_turn_by_beds'])); fu = H.escape(str(s['furnishing_cost_by_beds']))
    items.append(bi(f"Airbnb-style stays: platform {s['platform_fee_pct']:.0%}, average stay {s['avg_stay_nights']} nights, cleaning by bedrooms {cl}, supplies ${s['supplies_per_booked_night']}/night, insurance ${s['insurance_monthly']}/mo, furniture {fu} over {s['furnishing_amortization_months']} months, repairs {s['repairs_pct_of_revenue']:.0%}",
                    f"Estadías tipo Airbnb: plataforma {s['platform_fee_pct']:.0%}, estadía promedio {s['avg_stay_nights']} noches, limpieza por habitaciones {cl}, suministros ${s['supplies_per_booked_night']}/noche, seguro ${s['insurance_monthly']}/mes, muebles {fu} en {s['furnishing_amortization_months']} meses, reparaciones {s['repairs_pct_of_revenue']:.0%}"))
    items.append(bi(f"30+ night stays: occupancy {A['mtr']['occupancy']:.0%} (estimate); long-term tenant: vacancy {A['ownership_costs']['vacancy_pct_ltr']:.0%}",
                    f"Estadías de 30+ noches: ocupación {A['mtr']['occupancy']:.0%} (estimado); inquilino a largo plazo: desocupación {A['ownership_costs']['vacancy_pct_ltr']:.0%}"))
    return f"<h2>{bk('assump')}</h2><ul class='small'>" + "".join(f"<li>{i}</li>" for i in items) + "</ul>"


def _sources_html(r):
    seen, rows = set(), []
    nice_es = {"geocode": "búsqueda en el mapa", "listing_page": "página del anuncio", "fred": "tasas de Freddie Mac", "census": "Censo de EE.UU.",
               "insideairbnb": "Inside Airbnb", "craigslist": "Craigslist", "rent.com": "Rent.com", "redfin": "Redfin", "rentcast": "RentCast", "hud": "HUD"}
    for s in r.get("sources_status", []):
        if _lang() == "es":
            k0 = re.split(r"[:\s/]", str(s["source"]).strip().lower(), maxsplit=1)[0]
            s = dict(s, source=nice_es.get(k0, s["source"]))
        k = (s["source"], s["ok"])
        if k in seen:
            continue
        seen.add(k)
        note = "" if _lang() == "es" else H.escape(s.get("note") or "")     # engine notes are English-only
        http_ = str(s.get("http") or "")
        if _lang() == "es":
            http_ = {"cache": "guardado", "skipped": "omitido", "error": "error"}.get(http_, http_)
        rows.append([H.escape(s["source"]), "✅" if s["ok"] else "❌", H.escape(http_) + " " + note, _src_link(s["url"])])
    return f"<h2>{bk('sources')}</h2>" + _table([bi("Source", "Fuente"), "OK", bi("Status", "Estado"), bi("Link", "Enlace")], rows)


def property_md(r, lang):
    L = lambda k: t(k, lang)
    f, fs = r["facts"], r["fact_sources"]
    m = [f"# {L('prop_title')}", f"**{r['address']}** — {L('generated')} {r['generated']} ET", "", f"## {L('bottom_line')}"]
    m += [f"- {a if lang == 'en' else b}" for a, b in property_bullets(r)]
    if r["warnings"]:
        from .simple import plain_warnings
        m += ["", f"**{L('warnings')}:**"] + [f"- {a if lang == 'en' else b}" for a, b in plain_warnings(r)]
    m += ["", f"## {L('facts')}", f"| {L('field')} | {L('value')} | {L('source')} |", "|---|---|---|"]
    for k in ["price", "beds", "baths", "sqft", "hoa_monthly", "taxes_annual", "ownership", "hoa_includes"]:
        v = f.get(k)
        v = ", ".join(v) if isinstance(v, list) else v
        if k in ("price", "hoa_monthly", "taxes_annual") and v:
            v = money(float(v))
        m.append(f"| {L(k)} | {v if v not in (None, '') else '—'} | {fs.get(k, '')} |")
    sr = r["str_rules"]
    m += ["", f"## {L('legality')}", sr.get("summary_en" if lang == "en" else "summary_es", ""), "Sources: " + "; ".join(sr.get("sources", []))]
    le = r["ltr"]["estimate"]
    m += ["", f"## {L('ltr')}"]
    if le.get("ok"):
        m.append(f"**{money(le['median'])}**/mo ({money(le['p25'])}–{money(le['p75'])}, n={le['n']})")
        m += ["", "| Source | Listing | Rent | Bd | km |", "|---|---|---|---|---|"] + [f"| {c['source']} | {str(c.get('title'))[:40]} | {money(c['price'])} | {int(c['beds']) if c.get('beds') is not None else ''} | {c['dist_km']:.1f} |" for c in r["ltr"]["comps"][:8]]
    st = r["str"]
    if st.get("ok"):
        s = st["summary"]
        m += ["", f"## {L('str')}", f"n={s['n']} · ADR {money(s['adr_median'])} · occ {pct(s['occ_median_sf_model'])} · rev {money(s['revenue_median'])}/yr · datasets {', '.join(st['datasets'])}"]
    m += ["", f"## {L('scen')}", "| | " + " | ".join(s["label_en" if lang == "en" else "label_es"] for s in r["scenarios"]) + " |", "|---" * (len(r["scenarios"]) + 1) + "|"]
    m.append(f"| **{L('net')}** | " + " | ".join(money(s["net_monthly"]) if s.get("net_monthly") is not None else L("not_legal") for s in r["scenarios"]) + " |")
    m.append(f"| {L('total_cost')} | " + " | ".join(money(s.get("total_cost")) for s in r["scenarios"]) + " |")
    m += ["", f"_{L('disclaimer')}_"]
    return "\n".join(m)


# ------------------------------------------------------------------ arbitrage
def arb_bullets(r):
    sr, s = r["str_rules"], r["summary"]
    out = []
    if not sr.get("str_legal_for_tenant"):
        out.append((f"LEGALITY: tenant-operated Airbnb (rental arbitrage) is NOT legal in {r['town']} ({sr.get('status')}). The STR ranking below is for education only. {sr.get('summary_en','')}",
                    f"LEGALIDAD: el Airbnb operado por inquilinos (arbitraje) NO es legal en {r['town']} ({sr.get('status')}). La clasificación STR es solo educativa. {sr.get('summary_es','')}"))
    out.append((f"Scored {r['n_listings_scored']} of {r['n_listings_fetched']} fetched 1–3BR rentals. Median asking rent {money(s['median_rent'])}. Median estimated STR profit {money(s['median_str_profit'])}/mo; {pct(s['share_str_profitable'])} show a positive STR profit.",
                f"Se evaluaron {r['n_listings_scored']} de {r['n_listings_fetched']} alquileres de 1–3 hab. Renta mediana {money(s['median_rent'])}. Ganancia STR estimada mediana {money(s['median_str_profit'])}/mes; {pct(s['share_str_profitable'])} con ganancia positiva."))
    out.append((f"Legal alternative (30+ night furnished rentals, landlord permission required): median {money(s['median_mtr_profit'])}/mo; {pct(s['share_mtr_profitable'])} positive.",
                f"Alternativa legal (amueblado 30+ noches, con permiso del dueño): mediana {money(s['median_mtr_profit'])}/mes; {pct(s['share_mtr_profitable'])} positivos."))
    if not r.get("iab_covers_town"):
        out.append((f"No Inside Airbnb dataset covers {r['town']}; STR estimates use nearby datasets {', '.join(r['iab_datasets'])} as a proxy.",
                    f"No hay datos de Inside Airbnb para {r['town']}; se usan datos cercanos como aproximación."))
    return out


ARB_COLS = ["rank", "source", "title", "beds", "rent", "adr_median", "occ_sf_model", "str_revenue_annual", "str_profit_monthly", "mtr_profit_monthly", "legality_flag", "check_flag", "url"]


@with_lang
def arb_html(r, top=25, lang="both"):
    parts = [_open(f"BellaZu Rental Finder: {H.escape(r['town'])}", lang),
             f"<h1>{bk('arb_title')}: {H.escape(r['town'])}, {r['state'].upper()}</h1><div class='small'>{bk('generated')}: {r['generated']} ET · BellaZu v0.1</div>",
             f"<h2>{bk('bottom_line')}</h2><ul>" + "".join(f"<li>{bi(H.escape(a), H.escape(b), 'div')}</li>" for a, b in arb_bullets(r)) + "</ul>"]
    sr = r["str_rules"]
    parts.append(f"<div class='box {'bad' if not sr.get('str_legal_for_tenant') else 'ok'}'><b>{bk('legality')}</b><br>{bi(H.escape(sr.get('summary_en','')), H.escape(sr.get('summary_es','')))}<div class='small'>" + "; ".join(_rule_src(x) for x in sr.get("sources", [])) + f" · {bi('verified', 'verificado')} {sr.get('last_verified')}</div></div>")
    rows = [[b["beds"], b["n"], money(b["rent"]), money(b["str_rev"]), money(b["str_profit"]), money(b["mtr_profit"])] for b in r["summary"]["by_beds"]]
    parts.append(f"<h2>{bk('by_beds')}</h2>" + _table(["Bd", "n", "Median rent", "STR rev/yr", "STR profit/mo", "30+night profit/mo"], rows, (0, 1, 2, 3, 4, 5)))
    rows = []
    for x in r["results"][:top]:
        rows.append([x["rank"], H.escape(x["source"]), H.escape(str(x["title"]))[:40] + (" ⚠️" if x.get("check_flag") else ""), x["beds"], money(x["rent"]), money(x["adr_median"]), pct(x["occ_sf_model"]),
                     money(x["str_revenue_annual"]), f"<span class='{'neg' if x['str_profit_monthly'] < 0 else 'pos'}'>{money(x['str_profit_monthly'])}</span>",
                     f"<span class='{'neg' if x['mtr_profit_monthly'] < 0 else 'pos'}'>{money(x['mtr_profit_monthly'])}</span>", f"<a href='{safe_url(x.get('url'))}'>link</a>"])
    sb = r.get("sorted_by")
    rk = bi("Rentals ranked by estimated 30+ night (furnished, mid-term) profit — the legal option here", "Alquileres ordenados por ganancia estimada de 30+ noches (amueblado, mediano plazo) — la opción legal aquí") if sb == "mtr_profit_monthly" else bk("ranked")
    parts.append(f"<h2>{rk} (top {top})</h2>" + _table(["#", "Src", "Listing", "Bd", "Rent", "ADR", "Occ", "STR rev/yr", "STR profit/mo", "30+n profit/mo", ""], rows, (0, 3, 4, 5, 6, 7, 8, 9)))
    parts.append("<p class='small'>" + bi("STR profit = STR revenue/12 − rent − platform fees − cleaning − supplies − utilities − STR insurance − furniture (amortized) − repairs. ⚠️ = rent far below area median: verify listing (possible scam, room-only, or typo). Full list in the XLSX/CSV.",
                                          "Ganancia STR = ingreso STR/12 − renta − comisiones − limpieza − suministros − servicios − seguro STR − muebles (amortizados) − reparaciones. ⚠️ = renta muy por debajo de la mediana: verificar anuncio. Lista completa en XLSX/CSV.") + "</p>")
    parts.append(f"<h2>{bk('caveats')}</h2><ul class='small'>" + "".join(f"<li>{bi(H.escape(a), H.escape(b))}</li>" for a, b in zip(r["caveats_en"], r["caveats_es"])) + "</ul>")
    parts.append(_assump_html(r))
    parts.append(_sources_html(r))
    parts.append(f"<p class='small'>{bk('disclaimer')}</p>{_js()}</body></html>")
    return plainify("\n".join(parts), lang)


def arb_md(r, lang, top=15):
    L = lambda k: t(k, lang)
    m = [f"# {L('arb_title')}: {r['town']}, {r['state'].upper()}", f"{L('generated')} {r['generated']} ET", "", f"## {L('bottom_line')}"]
    m += [f"- {a if lang == 'en' else b}" for a, b in arb_bullets(r)]
    m += ["", f"## {L('ranked')}", "| # | Listing | Bd | Rent | STR rev/yr | STR profit/mo | 30+n profit/mo |", "|---|---|---|---|---|---|---|"]
    for x in r["results"][:top]:
        m.append(f"| {x['rank']} | [{str(x['title'])[:35]}]({x['url']}) | {x['beds']} | {money(x['rent'])} | {money(x['str_revenue_annual'])} | {money(x['str_profit_monthly'])} | {money(x['mtr_profit_monthly'])} |")
    m += ["", f"_{L('disclaimer')}_"]
    return "\n".join(m)


# ------------------------------------------------------------------ writers
def _flat_assumptions(A):
    rows = []
    for sect, vals in A.items():
        for k, v in vals.items():
            rows.append({"section": sect, "key": k, "value": json.dumps(v) if isinstance(v, (dict, list)) else v})
    return pd.DataFrame(rows)


SHEETS_ES = {"Compare": "Comparar", "Facts": "Datos", "Monthly numbers": "Números del mes", "Rentals nearby": "Alquileres cerca",
             "Rent estimates": "Estimados de renta", "RentCast examples": "Ejemplos RentCast", "Rooms nearby": "Cuartos cerca",
             "Airbnb nearby": "Airbnb cerca", "30+ night nearby": "30+ noches cerca", "Reference rents": "Rentas de referencia",
             "Airbnb rules": "Reglas de Airbnb", "Data sources": "Fuentes de datos", "Our estimates": "Nuestros estimados",
             "Rent by size": "Renta por tamaño", "Rooms": "Cuartos", "Airbnb market": "Mercado Airbnb", "Data log": "Registro de datos",
             "Rentals ranked": "Alquileres ordenados", "By bedrooms": "Por habitaciones", "Sources used": "Fuentes usadas"}
COLS_ES = {"option": "Opción", "label": "Etiqueta", "per_month": "Al mes", "detail": "Detalle", "rules": "Reglas", "field": "Dato", "value": "Valor",
           "source": "Fuente", "scenario": "Escenario", "legal": "Legal", "net_monthly": "Queda al mes", "total_cost": "Costo total al mes",
           "title": "Anuncio", "price": "Precio", "beds": "Habitaciones", "baths": "Baños", "sqft": "Pies cuadrados", "lat": "Latitud", "lon": "Longitud",
           "locality": "Localidad", "geo_locality": "Localidad (mapa)", "address": "Dirección", "url": "Enlace", "kind": "Tipo", "price_max": "Precio máximo",
           "beds_range": "Rango de habitaciones", "dist_km": "Distancia (km)", "name": "Nombre", "neighbourhood_cleansed": "Zona", "bedrooms": "Habitaciones",
           "price_num": "Precio por noche", "estimated_occupancy_l365d": "Noches ocupadas (12 meses)", "estimated_revenue_l365d": "Ingreso estimado (12 meses)",
           "number_of_reviews_ltm": "Reseñas (12 meses)", "minimum_nights": "Noches mínimas", "license": "Licencia", "listing_url": "Enlace", "dataset": "Datos de",
           "ok": "OK", "http": "Código HTTP", "note": "Nota", "benchmark": "Referencia", "zip": "Código postal", "area": "Área", "all": "Todas",
           "0br": "Estudio", "1br": "1 hab", "2br": "2 hab", "3br": "3 hab", "4br": "4 hab", "5br": "5 hab", "release": "Edición", "section": "Sección", "key": "Clave",
           "median": "Típica", "p25": "Parte baja del rango", "p75": "Parte alta del rango", "n": "Cantidad", "radius_km": "Radio (km)", "hud_safmr": "HUD SAFMR",
           "adr_median": "Precio por noche típico", "occ_median_sf_model": "Ocupación típica", "revenue_median": "Ingreso anual típico",
           "revenue_p25": "Ingreso anual (bajo)", "revenue_p75": "Ingreso anual (alto)", "days_old": "Días", "rent": "Renta", "rank": "Puesto",
           "free_median": "Comparables gratuitos (típica)", "free_n": "Comparables gratuitos (cantidad)", "rentcast_rent": "Renta RentCast", "rentcast_low": "RentCast bajo",
           "rentcast_high": "RentCast alto", "rentcast_n": "Comparables RentCast", "rentcast_ok": "RentCast OK", "rentcast_status": "Estado RentCast",
           "chosen": "Usado", "gap_pct": "Diferencia", "big_gap": "Gran diferencia", "down_payment": "Enganche", "loan_amount": "Monto del préstamo",
           "rate_pct": "Tasa (%)", "closing_costs_est": "Costos de cierre (estimado)", "cash_to_close_est": "Dinero para cerrar (estimado)",
           "roommate_rent": "Renta de compañeros", "rent_after_vacancy": "Renta tras desocupación", "status": "Estado", "min_nights_allowed": "Noches mínimas permitidas",
           "owner_str": "Airbnb del dueño", "tenant_str": "Airbnb del inquilino", "summary": "Resumen", "sources": "Fuentes", "source_type": "Tipo de fuente",
           "last_verified": "Verificado", "checked": "Revisado", "confidence": "Confianza", "str_legal_for_owner": "Airbnb legal para el dueño",
           "str_legal_for_tenant": "Airbnb legal para el inquilino", "hoa_monthly": "HOA / mantenimiento (mensual)", "ownership": "Tipo de vivienda",
           "hoa_includes": "El mantenimiento incluye", "year_built": "Año de construcción", "taxes_annual": "Impuestos (anual)", "description": "Descripción",
           "financing": "Financiamiento", "ownership_costs": "Costos de ser dueño", "roommate": "Compañeros", "str": "Airbnb", "mtr": "30+ noches",
           "comps": "Comparables", "arbitrage": "Arbitraje", "adr_p25": "Precio por noche (bajo)", "adr_p75": "Precio por noche (alto)",
           "occ_p25": "Ocupación (baja)", "occ_p75": "Ocupación (alta)", "cal_unavail_median": "Días ocupados en calendario (típico)",
           "licensed_share": "Parte con licencia", "free_ok": "Comparables gratuitos OK", "picture_url": "Foto (enlace)", "str_types": "Tipos de vivienda permitidos"}
VAL_ES = {"owner_occupied_permit_only": "solo dueño que vive allí, con permiso", "banned": "prohibido", "unknown": "no se sabe", "allowed": "permitido",
          "permit_required": "con permiso", "high": "alta", "medium": "media", "low": "baja", "official ordinance": "ordenanza oficial",
          "official code": "código oficial", "news": "noticias", "free": "gratuitos", "rentcast": "RentCast", "hud_safmr": "HUD SAFMR",
          "single-family": "casa unifamiliar", "two-family": "casa de dos familias", "multi-family": "multifamiliar", "townhouse": "casa adosada",
          "condo": "condominio", "co-op": "co-op (cooperativa)", "FOR_SALE": "en venta", "FOR_RENT": "en alquiler", "user input": "escrito por usted",
          "parsed from listing description": "tomado de la descripción del anuncio", "cache": "guardado"}
RULE_ROWS_EN_ONLY = {"owner_str", "tenant_str"}      # free English text; the Spanish summary row says the same


def _col_es(c):
    c = str(c)
    for pre, lab in (("cost_", "Costo: "), ("income_", "Ingreso: "), ("loan_", "Préstamo: ")):
        if c.startswith(pre):
            k = c[len(pre):]
            return lab + (COST[k][1] if k in COST else COLS_ES.get(k, k.replace("_", " ")))
    return COLS_ES.get(c) or (T_ES(c) or c.replace("_", " "))


def T_ES(k):
    v = t(k, "es")
    return v if v != k else None


def _one_lang(df, lang):
    """Keep only the chosen language: drop *_en/*_es columns (and field rows) of the other language, strip the suffix of the kept one."""
    if lang not in ("en", "es") or df is None or df.empty:
        return df
    other = "_es" if lang == "en" else "_en"
    df = df[[c for c in df.columns if not str(c).endswith(other)]]
    df = df.rename(columns={c: str(c)[:-3] for c in df.columns if str(c).endswith("_" + lang)})
    if "field" in df.columns:
        df = df[~df["field"].astype(str).str.endswith(other)].copy()
        if lang == "es":
            df = df[~df["field"].astype(str).isin(RULE_ROWS_EN_ONLY)].copy()
            if "value" in df.columns:
                df["value"] = df["value"].map(lambda v: ", ".join(VAL_ES.get(x.strip(" '"), x.strip(" '")) for x in v.strip("[]").split(",")) if isinstance(v, str) and v.startswith("[") and "http" not in v else v)
        df["field"] = df["field"].astype(str).map(lambda k: k[:-3] if k.endswith("_" + lang) else k)
        df["field"] = df["field"].map(lambda k: t(k, lang) if (T_ES(k) or t(k, "en") != k) else (COLS_ES.get(k, k) if lang == "es" else k.replace("_", " ")))
    if lang == "es" and "section" in df.columns:
        df = df.copy(); df["section"] = df["section"].map(lambda v: COLS_ES.get(v, v))
    return df


def _sheet(lang, name):
    return SHEETS_ES.get(name, name) if lang == "es" else name


def _style_xlsx(path, lang="en"):
    """Report look for the spreadsheet: ink header row with rose text, frozen header, readable widths, plain wording (values unchanged)."""
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = load_workbook(path)
    head_fill, head_font = PatternFill("solid", fgColor="141414"), Font(bold=True, color="FFD3DE", name="Calibri")
    for ws in wb.worksheets:
        ws.sheet_properties.tabColor = "D8567C"
        ws.freeze_panes = "A2"
        for row in ws.iter_rows():
            for c in row:
                if lang == "es" and isinstance(c.value, bool):
                    c.value = "Sí" if c.value else "No"
                elif lang == "es" and isinstance(c.value, str) and c.value in VAL_ES:
                    c.value = VAL_ES[c.value]
                elif lang == "es" and isinstance(c.value, str) and c.value.startswith("official code ("):
                    c.value = "código oficial (sin sección sobre Airbnb)"
                elif lang == "es" and isinstance(c.value, str) and c.value in ("True", "False"):
                    c.value = "Sí" if c.value == "True" else "No"
                if isinstance(c.value, str):
                    c.value = _plain_text(c.value, "es" if lang == "es" else "en")
        for c in ws[1]:
            c.fill, c.font = head_fill, head_font
            c.alignment = Alignment(vertical="center", wrap_text=True)
            if isinstance(c.value, str):
                c.value = _col_es(c.value) if lang == "es" else c.value.replace("_", " ").capitalize()
        for col in ws.columns:
            w = max((len(str(c.value)) for c in col[:200] if c.value is not None), default=8)
            ws.column_dimensions[col[0].column_letter].width = min(max(10, w + 2), 60)
    wb.save(path)


def write_property(r, outdir, cv=None, lang="both"):
    outdir = pathlib.Path(outdir); outdir.mkdir(parents=True, exist_ok=True)
    one = lang in ("en", "es")
    base = outdir / (f"BellaZu_Reporte_Casa_{slug(r['address'])}_{r['generated'][:10]}" if lang == "es" else f"BellaZu_Property_Report_{slug(r['address'])}_{r['generated'][:10]}")
    files = {}
    (p := base.with_suffix(".html")).write_text(property_html(r, lang=lang, cv=cv)); files["html"] = str(p)
    (p := base.with_suffix(".md")).write_text(property_md(r, lang) if one else property_md(r, "en") + "\n\n---\n\n" + property_md(r, "es")); files["md"] = str(p)
    (p := base.with_suffix(".json")).write_text(json.dumps(r, indent=1, default=str)); files["json"] = str(p)
    sc_rows = []
    for s in r["scenarios"]:
        row = ({"scenario": s["label_es" if lang == "es" else "label_en"]} if one else {"scenario_en": s["label_en"], "scenario_es": s["label_es"]})
        row.update({"legal": s.get("legal"), "net_monthly": s.get("net_monthly"), "total_cost": s.get("total_cost")})
        row.update({f"cost_{k}": v for k, v in (s.get("costs") or {}).items()})
        row.update({f"income_{k}": v for k, v in (s.get("income") or {}).items()})
        if s.get("loan"): row.update({f"loan_{k}": v for k, v in s["loan"].items()})
        sc_rows.append(row)
    sc = pd.DataFrame(sc_rows)
    (p := base.with_name(base.name + "_scenarios.csv")); sc.to_csv(p, index=False); files["csv"] = str(p)
    p = base.with_suffix(".xlsx")
    X = lambda df, name: _one_lang(pd.DataFrame(df) if not isinstance(df, pd.DataFrame) else df, lang).to_excel(w, sheet_name=_sheet(lang, name), index=False)
    with pd.ExcelWriter(p, engine="openpyxl") as w:
        facts = pd.DataFrame([{"field": k, "value": (", ".join(v) if isinstance(v, list) else v), "source": r["fact_sources"].get(k, "")} for k, v in r["facts"].items() if k != "description"])
        if cv:
            X(_cv_rows(cv, lang), "Compare")
        X(facts, "Facts")
        X(sc, "Monthly numbers")
        X(r["ltr"]["comps"], "Rentals nearby")
        X([{k: v for k, v in (r.get("rent_compare") or {}).items()}], "Rent estimates")
        X((r.get("rentcast") or {}).get("avm_comps") or [], "RentCast examples")
        X(r["rooms"].get("comps", []), "Rooms nearby")
        X(r["str"].get("comps", []), "Airbnb nearby")
        X(r["mtr"].get("comps", []), "30+ night nearby")
        bm = r["benchmarks"]
        X([{"benchmark": k, **(v if isinstance(v, dict) else {})} for k, v in bm.items() if v], "Reference rents")
        X([{"field": k, "value": str(v)} for k, v in r["str_rules"].items()], "Airbnb rules")
        X([{k: v for k, v in x.items() if not (lang == "es" and k == "note")} for x in r["sources_status"]], "Data sources")
        X(_flat_assumptions(r["assumptions"]), "Our estimates")
    _style_xlsx(p, lang)
    files["xlsx"] = str(p)
    return files


def write_arbitrage(r, outdir):
    outdir = pathlib.Path(outdir); outdir.mkdir(parents=True, exist_ok=True)
    base = outdir / f"BellaZu_Arbitrage_Scan_{slug(r['town'])}_{r['state'].upper()}_{r['generated'][:10]}"
    files = {}
    (p := base.with_suffix(".html")).write_text(arb_html(r)); files["html"] = str(p)
    (p := base.with_suffix(".md")).write_text(arb_md(r, "en") + "\n\n---\n\n" + arb_md(r, "es")); files["md"] = str(p)
    (p := base.with_suffix(".json")).write_text(json.dumps(r, indent=1, default=str)); files["json"] = str(p)
    df = pd.DataFrame(r["results"])
    (p := base.with_suffix(".csv")); df.to_csv(p, index=False); files["csv"] = str(p)
    p = base.with_suffix(".xlsx")
    with pd.ExcelWriter(p, engine="openpyxl") as w:
        df.to_excel(w, sheet_name="Rentals ranked", index=False)
        pd.DataFrame(r["summary"]["by_beds"]).to_excel(w, sheet_name="By bedrooms", index=False)
        pd.DataFrame([{"field": k, "value": str(v)} for k, v in r["str_rules"].items()]).to_excel(w, sheet_name="Airbnb rules", index=False)
        pd.DataFrame(r["sources_used"]).to_excel(w, sheet_name="Sources used", index=False)
        pd.DataFrame(r["sources_status"]).to_excel(w, sheet_name="Data log", index=False)
        _flat_assumptions(r["assumptions"]).to_excel(w, sheet_name="Our estimates", index=False)
    _style_xlsx(p)
    files["xlsx"] = str(p)
    return files


# ------------------------------------------------------------------ town snapshot report
@with_lang
def town_html(a, cv=None, lang="both"):
    sr = a.get("str_rules") or {}
    mk = (a.get("market") or {}).get("summary") or {}
    d = a.get("drive") or {}
    parts = [_open(("Reporte del pueblo BellaZu" if lang == "es" else "BellaZu Town Report") + f": {H.escape(a['town'])}", lang),
             f"<h1>{H.escape(a['town'])}</h1><div class='small'>{bk('generated')}: {a['generated']} ET · BellaZu v0.1</div>"]
    parts.append(compare_html(cv) if cv else "<p>" + bi("Tap a price in the app to add the monthly compare.", "Toque un precio en la app para agregar la comparación mensual.") + "</p>")
    cls = "bad" if not sr.get("str_legal_for_owner") else "ok"
    parts.append(f"<h2>{bk('legality')}</h2><div class='box {cls}'>{bi(H.escape(sr.get('summary_en', '')), H.escape(sr.get('summary_es', '')))}<div class='small'>" + bi("Sources", "Fuentes") + ": "
                 + "; ".join(_rule_src(x) for x in sr.get("sources", [])) + f" · {bi('checked', 'revisado')} {sr.get('checked') or sr.get('last_verified')} · {bi('confidence', 'confianza')} {H.escape(_conf(sr.get('confidence', '?')))}</div></div>")
    rows = []
    for b, v in sorted(((int(k), v) for k, v in (a.get("by_beds") or {}).items())):
        e = v.get("ltr") or {}
        rows.append([(f"{b} " + bi("bd", "hab")) if b else bi("Studio", "Estudio"), money(e.get("median")) if e.get("ok") else "—", f"{money(e.get('p25'))}–{money(e.get('p75'))}" if e.get("ok") else "—", e.get("n") or 0, money(v.get("hud"))])
    parts.append(f"<h2>{bi('Rent by size', 'Renta por tamaño')}</h2>" + _table([bi("Size", "Tamaño"), bi("Typical", "Típica"), bi("Usual range", "Rango usual"), "n", "HUD SAFMR"], rows, (1, 3, 4)))
    ro = a.get("rooms") or {}
    if ro.get("ok"):
        parts.append(f"<p>{bi('Room in a shared home', 'Cuarto en casa compartida')}: <b>{money(ro['median'])}</b>{_mo()} ({money(ro['p25'])}–{money(ro['p75'])}, n={ro['n']}, Craigslist)</p>")
    if mk.get("n"):
        rough = "" if (a.get("market") or {}).get("covered") else bi(" Rough estimate borrowed from nearby city data.", " Estimado aproximado con datos de una ciudad cercana.")
        parts.append(f"<h2>{bi('Airbnb market nearby', 'Mercado de Airbnb cerca')}</h2><p>n={mk['n']} · {bi('typical nightly', 'noche típica')} <b>{money(mk.get('adr_median'))}</b> · "
                     f"{bi('nights/yr', 'noches/año')} <b>{round((mk.get('occ_median_sf_model') or 0) * 365)}</b> · {bi('income/yr', 'ingreso/año')} <b>{money(mk.get('revenue_median'))}</b> "
                     f"({money(mk.get('revenue_p25'))}–{money(mk.get('revenue_p75'))}).{rough}</p>")
    s = a.get("seasonality") or {}
    if (s.get("short") or {}).get("index"):
        mn = (["Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Sep", "Oct", "Nov", "Dic"] if lang == "es" else
              ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"])
        parts.append(f"<h2>{bi('Busy vs slow months', 'Meses de mucho y poco movimiento')}</h2>" + _table(mn, [[f"{v:.2f}" for v in s["short"]["index"]]], tuple(range(12)))
                     + "<p class='small'>" + bi("1.00 = average month. Guest reviews per month (a stand-in for bookings), Inside Airbnb, " + H.escape(str(s["city"]).replace("-", " ").title()), "1.00 = mes promedio. Reseñas por mes (aproximación de reservas), Inside Airbnb, " + H.escape(str(s["city"]).replace("-", " ").title())) + "</p>")
    if d:
        r0, r1 = d.get('rush', ['?', '?'])[0], d.get('rush', ['?', '?'])[1]
        parts.append("<p class='small'>" + bi(f"Drive to Midtown: about {d.get('min')} min off-peak, {r0}-{r1} min at rush hour via {H.escape(str(d.get('crossing')))}. Typical, not live.",
                                              f"En carro a Midtown: unos {d.get('min')} min sin tráfico, {r0}-{r1} min en hora pico por {H.escape(str(d.get('crossing')))}. Típico, no en vivo.") + "</p>")
    parts.append(f"<p class='small'>{bk('disclaimer')}</p>{_js()}</body></html>")
    return plainify("\n".join(parts), lang)


def _conf(c):
    return {"high": "alta", "medium": "media", "low": "baja"}.get(str(c).lower(), str(c)) if _lang() == "es" else str(c)


def write_town(a, cv, outdir, lang="both"):
    outdir = pathlib.Path(outdir); outdir.mkdir(parents=True, exist_ok=True)
    p = outdir / (f"BellaZu_Pueblo_{slug(a['town'])}_{a['generated'][:10]}.xlsx" if lang == "es" else f"BellaZu_Town_{slug(a['town'])}_{a['generated'][:10]}.xlsx")
    X = lambda df, name: _one_lang(pd.DataFrame(df) if not isinstance(df, pd.DataFrame) else df, lang).to_excel(w, sheet_name=_sheet(lang, name)[:31], index=False)
    with pd.ExcelWriter(p, engine="openpyxl") as w:
        if cv:
            X(_cv_rows(cv, lang), "Compare")
        X([{"beds": k, **{kk: vv for kk, vv in (v.get("ltr") or {}).items()}, "hud_safmr": v.get("hud")} for k, v in (a.get("by_beds") or {}).items()], "Rent by size")
        for b, v in (a.get("by_beds") or {}).items():
            if v.get("ltr_comps"):
                X(v["ltr_comps"], f"Alquileres {b} hab" if lang == "es" else f"Rentals {b}bd")
        X((a.get("rooms") or {}).get("comps") or [], "Rooms")
        X([(a.get("market") or {}).get("summary") or {}], "Airbnb market")
        X([{"field": k, "value": str(v)} for k, v in (a.get("str_rules") or {}).items()], "Airbnb rules")
        X([{k: v for k, v in x.items() if not (lang == "es" and k == "note")} for x in (a.get("sources_status") or [])], "Data log")
        X(_flat_assumptions(a["assumptions"]), "Our estimates")
    _style_xlsx(p, lang)
    return str(p)
